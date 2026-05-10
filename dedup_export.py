"""
Deduplicate the scraped directory and export an XLSX workbook with the
member photos embedded as images.

Reads:  data/concordia_directory.csv
Writes: data/concordia_directory_deduped.csv
        data/concordia_directory.xlsx
        data/photos/*.{png,jpg,jpeg}    (download cache)

Dedup rule: rows are grouped by lower-cased full name. For every field the
longest non-empty value across the group wins (so the more complete profile
beats the sparser one).
"""

from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlparse

import requests
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = Path(__file__).parent
INPUT_CSV = ROOT / "data" / "concordia_directory.csv"
DEDUPED_CSV = ROOT / "data" / "concordia_directory_deduped.csv"
OUTPUT_XLSX = ROOT / "data" / "concordia_directory.xlsx"
PHOTO_CACHE = ROOT / "data" / "photos"

FIELDS = [
    "full_name", "title", "organization", "bio", "industry",
    "city", "state", "country", "website", "region_of_interest",
    "twitter", "linkedin", "photo",
]

PHOTO_PX = 80          # rendered side, pixels (square)
DOWNLOAD_WORKERS = 12

# Excel forbids these control characters in cell values (XML 1.0 spec).
_ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(v: str) -> str:
    if not v:
        return ""
    return _ILLEGAL_CHARS.sub("", v)


def read_rows() -> list[dict]:
    with INPUT_CSV.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def dedupe(rows: list[dict]) -> list[dict]:
    groups: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        groups[r["full_name"].strip().lower()].append(r)

    merged: list[dict] = []
    for items in groups.values():
        if len(items) == 1:
            merged.append(items[0])
            continue
        out = {f: "" for f in FIELDS}
        for f in FIELDS:
            best = ""
            for it in items:
                v = (it.get(f) or "").strip()
                if len(v) > len(best):
                    best = v
            out[f] = best
        merged.append(out)
    merged.sort(key=lambda r: r["full_name"].lower())
    return merged


def write_deduped_csv(rows: list[dict]) -> None:
    with DEDUPED_CSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, quoting=csv.QUOTE_ALL)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in FIELDS})


def cache_path_for(url: str) -> Path:
    return PHOTO_CACHE / Path(urlparse(url).path).name


def thumb_path_for(cached: Path) -> Path:
    return cached.with_name(cached.stem + ".thumb.png")


def download(url: str) -> Path | None:
    target = cache_path_for(url)
    if target.exists() and target.stat().st_size > 0:
        return target
    try:
        r = requests.get(url, timeout=30)
        r.raise_for_status()
        target.write_bytes(r.content)
        return target
    except Exception as e:
        print(f"  download failed: {url}: {e}", file=sys.stderr)
        return None


def make_thumb(cached: Path) -> Path | None:
    thumb = thumb_path_for(cached)
    if thumb.exists() and thumb.stat().st_size > 0:
        return thumb
    try:
        with PILImage.open(cached) as im:
            im = im.convert("RGB")
            side = min(im.size)
            left = (im.width - side) // 2
            top = (im.height - side) // 2
            im = im.crop((left, top, left + side, top + side))
            im = im.resize((PHOTO_PX, PHOTO_PX), PILImage.LANCZOS)
            im.save(thumb, "PNG", optimize=True)
        return thumb
    except Exception as e:
        print(f"  thumb failed: {cached.name}: {e}", file=sys.stderr)
        return None


def fetch_all_photos(rows: list[dict]) -> dict[str, Path]:
    """Download every unique photo URL and produce its square thumbnail.
    Returns map: url -> thumbnail path (or absent on failure)."""
    PHOTO_CACHE.mkdir(parents=True, exist_ok=True)
    urls = sorted({r["photo"].strip() for r in rows if r.get("photo", "").strip()})
    print(f"downloading {len(urls)} unique photos with {DOWNLOAD_WORKERS} workers...")

    cached: dict[str, Path] = {}
    with ThreadPoolExecutor(max_workers=DOWNLOAD_WORKERS) as ex:
        futures = {ex.submit(download, u): u for u in urls}
        for i, fut in enumerate(as_completed(futures), start=1):
            url = futures[fut]
            path = fut.result()
            if path:
                cached[url] = path
            if i % 100 == 0 or i == len(urls):
                print(f"  downloaded {i}/{len(urls)}")

    print("generating thumbnails...")
    thumbs: dict[str, Path] = {}
    for i, (url, path) in enumerate(cached.items(), start=1):
        thumb = make_thumb(path)
        if thumb:
            thumbs[url] = thumb
        if i % 200 == 0 or i == len(cached):
            print(f"  thumbed {i}/{len(cached)}")
    return thumbs


def write_xlsx(rows: list[dict], thumbs: dict[str, Path]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Directory"

    headers = ["photo"] + [f for f in FIELDS if f != "photo"] + ["photo_url"]
    ws.append(headers)

    widths = {
        "photo": 14, "full_name": 26, "title": 32, "organization": 30,
        "bio": 70, "industry": 22, "city": 18, "state": 14,
        "country": 18, "website": 32, "region_of_interest": 24,
        "twitter": 22, "linkedin": 38, "photo_url": 50,
    }
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 18)
    ws.row_dimensions[1].height = 22

    point_height = PHOTO_PX * 0.75 + 6   # row height in Excel points

    for idx, r in enumerate(rows, start=2):
        for col_i, h in enumerate(headers, start=1):
            if h == "photo":
                continue
            value = r.get("photo", "") if h == "photo_url" else r.get(h, "")
            ws.cell(row=idx, column=col_i, value=clean(value))
        ws.row_dimensions[idx].height = point_height

        url = r.get("photo", "").strip()
        thumb = thumbs.get(url)
        if not thumb:
            continue
        try:
            img = XLImage(str(thumb))
            ws.add_image(img, f"A{idx}")
        except Exception as e:
            print(f"  embed failed for {thumb.name}: {e}", file=sys.stderr)

        if idx % 200 == 0:
            print(f"  embedded {idx - 1}/{len(rows)}")

    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_XLSX)


def main() -> int:
    rows = read_rows()
    print(f"read {len(rows)} rows from {INPUT_CSV.name}")
    deduped = dedupe(rows)
    print(f"deduped to {len(deduped)} rows")
    write_deduped_csv(deduped)
    print(f"wrote {DEDUPED_CSV}")

    thumbs = fetch_all_photos(deduped)
    print(f"have thumbnails for {len(thumbs)} photos")

    write_xlsx(deduped, thumbs)
    print(f"wrote {OUTPUT_XLSX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
